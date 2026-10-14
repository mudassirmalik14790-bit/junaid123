"""AI Resume ATS Checker - Streamlit app powered by Google Gemini Flash."""

import io
import json
import os
import re
import time

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-3.8-flash"
MODEL_NAME = os.getenv("GEMINI_MODEL", DEFAULT_MODEL)  # override via env var or Streamlit secret
MAX_CHARS = 20000  # keep the prompt a sensible size
MIN_CHARS = 100  # below this, the file is probably a scanned image
RETRY_CODES = {429, 500, 502, 503, 504}  # temporary errors worth retrying
RETRY_DELAYS = (2, 5, 10)  # seconds to wait before each retry


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a PDF, DOCX or TXT file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many resumes keep content inside tables
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT.")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def build_prompt(resume_text: str, job_description: str = "") -> str:
    """Create the instruction prompt sent to Gemini."""
    jd_block = (
        f"\nTARGET JOB DESCRIPTION:\n\"\"\"\n{job_description.strip()[:8000]}\n\"\"\"\n"
        "Judge keyword match against this job description.\n"
        if job_description.strip()
        else "\nNo job description was provided. Judge against general best "
        "practices for the role the resume appears to target.\n"
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and resume coach.
Analyse the resume below and return ONLY a JSON object (no markdown) with this exact shape:

{{
  "ats_score": <integer 0-100>,
  "score_breakdown": {{
    "formatting": <integer 0-100>,
    "keywords": <integer 0-100>,
    "content_quality": <integer 0-100>,
    "structure": <integer 0-100>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<string>", ...],
  "weaknesses": ["<string>", ...],
  "missing_keywords": ["<string>", ...],
  "improvements": [
    {{"section": "<resume section>", "issue": "<what is wrong>",
      "suggestion": "<specific fix, with an example rewrite where useful>",
      "priority": "High|Medium|Low"}}
  ]
}}

Scoring guidance: check parseable formatting, standard section headings
(Summary, Experience, Education, Skills), contact details, quantified achievements,
strong action verbs, relevant keywords, consistent dates, and length.
Be honest and strict; do not inflate the score. Give 5-8 improvements, most important first.
{jd_block}
RESUME:
\"\"\"
{resume_text[:MAX_CHARS]}
\"\"\"
"""


def parse_response(raw: str) -> dict:
    """Parse and normalise the model's JSON output."""
    cleaned = re.sub(r"^```(?:json)?|```$", "", (raw or "").strip(), flags=re.M).strip()
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, re.S)  # fall back to the outermost object
        if not match:
            raise ValueError("The AI response was not valid JSON. Please try again.")
        data = json.loads(match.group(0))

    def clamp(value) -> int:
        try:
            return max(0, min(100, int(round(float(value)))))
        except (TypeError, ValueError):
            return 0

    breakdown = data.get("score_breakdown") or {}
    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            priority = str(item.get("priority", "Medium")).capitalize()
            improvements.append(
                {
                    "section": str(item.get("section", "General")),
                    "issue": str(item.get("issue", "")),
                    "suggestion": str(item.get("suggestion", "")),
                    "priority": priority if priority in ("High", "Medium", "Low") else "Medium",
                }
            )
    return {
        "ats_score": clamp(data.get("ats_score")),
        "score_breakdown": {
            k: clamp(breakdown.get(k))
            for k in ("formatting", "keywords", "content_quality", "structure")
        },
        "summary": str(data.get("summary", "")),
        "strengths": [str(s) for s in data.get("strengths") or []],
        "weaknesses": [str(s) for s in data.get("weaknesses") or []],
        "missing_keywords": [str(s) for s in data.get("missing_keywords") or []],
        "improvements": improvements,
    }


def analyze_resume(api_key: str, resume_text: str, job_description: str = "") -> dict:
    """Send the resume to Gemini (retrying on temporary overload) and parse the result."""
    client = genai.Client(api_key=api_key)
    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
    )
    prompt = build_prompt(resume_text, job_description)

    for attempt in range(len(RETRY_DELAYS) + 1):
        try:
            response = client.models.generate_content(
                model=get_model_name(), contents=prompt, config=config
            )
            return parse_response(response.text)
        except Exception as exc:
            code = getattr(exc, "code", None)
            if code in RETRY_CODES and attempt < len(RETRY_DELAYS):
                time.sleep(RETRY_DELAYS[attempt])
                continue
            if code in RETRY_CODES:
                raise RuntimeError(
                    "Google's Gemini servers are busy right now. "
                    "Please wait a minute and click Analyze again."
                ) from exc
            raise


def get_model_name() -> str:
    """Model from Streamlit secrets (GEMINI_MODEL), else env var, else the default."""
    try:
        if "GEMINI_MODEL" in st.secrets:
            return st.secrets["GEMINI_MODEL"]
    except Exception:
        pass
    return MODEL_NAME


def get_api_key() -> str:
    """Look for the key in Streamlit secrets, then env vars, then the sidebar."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:  # no secrets file present
        pass
    env_key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if env_key:
        return env_key
    return st.sidebar.text_input(
        "Gemini API key",
        type="password",
        help="Get a free key at https://aistudio.google.com/apikey",
    )


def score_label(score: int) -> str:
    if score >= 80:
        return "🟢 Excellent"
    if score >= 60:
        return "🟡 Good, needs work"
    return "🔴 Needs major improvement"


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="centered")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it.")

    api_key = get_api_key()

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste the job posting here for a keyword match against that role...",
    )

    if st.button("Analyze resume", type="primary", use_container_width=True):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return
        if uploaded is None:
            st.error("Please upload a resume first.")
            return

        try:
            resume_text = extract_text(uploaded.getvalue(), uploaded.name)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return
        if len(resume_text) < MIN_CHARS:
            st.error(
                "Very little text was found. The file may be a scanned image. "
                "Please upload a text-based PDF or a DOCX."
            )
            return

        with st.spinner("Analyzing your resume..."):
            try:
                result = analyze_resume(api_key, resume_text, job_description)
            except Exception as exc:
                st.error(f"Analysis failed: {exc}")
                return
        st.session_state["result"] = result

    result = st.session_state.get("result")
    if not result:
        return

    st.divider()
    score = result["ats_score"]
    col1, col2 = st.columns([1, 2])
    col1.metric("ATS Score", f"{score}/100")
    col2.markdown(f"### {score_label(score)}")
    col2.progress(score / 100)
    if result["summary"]:
        st.info(result["summary"])

    st.subheader("Score breakdown")
    cols = st.columns(4)
    for col, (name, value) in zip(cols, result["score_breakdown"].items()):
        col.metric(name.replace("_", " ").title(), f"{value}")

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for s in result["strengths"] or ["None listed"]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("⚠️ Weaknesses")
        for s in result["weaknesses"] or ["None listed"]:
            st.markdown(f"- {s}")

    if result["missing_keywords"]:
        st.subheader("🔑 Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.subheader("🛠️ Suggested improvements")
    icons = {"High": "🔴", "Medium": "🟡", "Low": "🟢"}
    order = {"High": 0, "Medium": 1, "Low": 2}
    for item in sorted(result["improvements"], key=lambda i: order[i["priority"]]):
        with st.expander(f"{icons[item['priority']]} {item['section']} ({item['priority']})"):
            st.markdown(f"**Issue:** {item['issue']}")
            st.markdown(f"**Fix:** {item['suggestion']}")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


if __name__ == "__main__":
    main()
